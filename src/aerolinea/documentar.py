"""Diccionario de datos y diagramas físicos generados desde el catálogo del motor.

Fuentes:
  - Catálogo de SQL Server (sys.*): tipos, longitudes, nulabilidad, valores
    predeterminados, columnas calculadas, PK, FK, UNIQUE, índices únicos
    filtrados, CHECK, enmascaramiento y número de filas.
  - docs/02_modelo/definiciones_atributos.csv: definición de negocio y dominio
    redactados por el equipo.
Así el diccionario y el diagrama físico coinciden con el código por construcción.
"""
from __future__ import annotations

import csv
import re
import shutil
import subprocess
from collections import defaultdict
from pathlib import Path

from .config import RAIZ_PROYECTO, ConfigBD
from .db import conectar, ejecutar

DIR_MODELO = RAIZ_PROYECTO / "docs" / "02_modelo"
DIR_DIAGRAMAS = DIR_MODELO / "diagramas"
ARCHIVO_DEFINICIONES = DIR_MODELO / "definiciones_atributos.csv"
SCRIPT_DESCRIPCIONES = RAIZ_PROYECTO / "sql" / "13_descripciones.sql"
ESQUEMAS_EXCLUIDOS = ("stg",)

AREAS = {
    "fisico_red_operaciones": ("Red, flota, vuelos y tripulación", [
        "ref.localizacion", "ref.aeropuerto", "ref.zona_horaria", "ref.estado_eeuu", "ops.ruta", "ops.programacion",
        "ops.programacion_dia", "ops.vuelo", "ops.vuelo_cabina", "ops.avion", "ops.configuracion_asientos",
        "ops.asiento_configuracion", "ref.modelo_avion", "ref.cabina", "ops.asignacion_tripulacion", "rh.empleado",
        "rh.cargo", "rh.departamento", "ref.estado_vuelo"]),
    "fisico_ventas": ("Personas, compras, tiquetes, tarifas, pagos y reembolsos", [
        "ven.persona", "ven.cliente", "ven.compra", "ven.tiquete", "ven.tarifa", "ven.descuento", "ven.pago",
        "ven.reembolso", "ven.politica_cancelacion", "ref.metodo_pago", "ref.estado_tiquete", "ref.estado_pago",
        "ref.motivo_cancelacion", "ref.pais", "ops.vuelo_cabina", "ops.ruta", "ref.localizacion", "rh.empleado", "ref.cabina"]),
    "fisico_catalogos_control": ("Catálogos de estado, transiciones, parámetros y control", [
        "ref.estado_vuelo", "ref.estado_tiquete", "ref.estado_pago", "ref.transicion_vuelo", "ref.transicion_tiquete",
        "ref.transicion_pago", "ref.parametro", "ref.mensaje_error", "aud.cambio_estado", "adm.registro_despliegue",
        "tst.resultado_prueba"]),
}

SQL_COLUMNAS = """
SELECT s.name AS esquema, t.name AS tabla, c.column_id, c.name AS columna, ty.name AS tipo, c.max_length, c.precision,
       c.scale, c.is_nullable, c.is_identity, c.is_computed, cc.definition AS calculo, cc.is_persisted,
       dc.definition AS predeterminado, mc.masking_function
FROM sys.tables AS t
JOIN sys.schemas AS s ON s.schema_id = t.schema_id
JOIN sys.columns AS c ON c.object_id = t.object_id
JOIN sys.types AS ty ON ty.user_type_id = c.user_type_id
LEFT JOIN sys.computed_columns AS cc ON cc.object_id = c.object_id AND cc.column_id = c.column_id
LEFT JOIN sys.default_constraints AS dc ON dc.parent_object_id = c.object_id AND dc.parent_column_id = c.column_id
LEFT JOIN sys.masked_columns AS mc ON mc.object_id = c.object_id AND mc.column_id = c.column_id
ORDER BY s.name, t.name, c.column_id;
"""
SQL_INDICES = """
SELECT s.name AS esquema, t.name AS tabla, i.name AS indice, i.is_primary_key, i.is_unique_constraint, i.is_unique,
       i.filter_definition, c.name AS columna, ic.key_ordinal
FROM sys.indexes AS i
JOIN sys.tables AS t ON t.object_id = i.object_id
JOIN sys.schemas AS s ON s.schema_id = t.schema_id
JOIN sys.index_columns AS ic ON ic.object_id = i.object_id AND ic.index_id = i.index_id AND ic.is_included_column = 0
JOIN sys.columns AS c ON c.object_id = ic.object_id AND c.column_id = ic.column_id
WHERE i.is_unique = 1 OR i.is_primary_key = 1
ORDER BY s.name, t.name, i.name, ic.key_ordinal;
"""
SQL_FK = """
SELECT fk.name AS fk, sp.name AS esquema, tp.name AS tabla, cp.name AS columna, sr.name AS esquema_ref, tr.name AS tabla_ref,
       cr.name AS columna_ref, cp.is_nullable, fkc.constraint_column_id
FROM sys.foreign_keys AS fk
JOIN sys.foreign_key_columns AS fkc ON fkc.constraint_object_id = fk.object_id
JOIN sys.tables AS tp ON tp.object_id = fk.parent_object_id
JOIN sys.schemas AS sp ON sp.schema_id = tp.schema_id
JOIN sys.columns AS cp ON cp.object_id = fkc.parent_object_id AND cp.column_id = fkc.parent_column_id
JOIN sys.tables AS tr ON tr.object_id = fk.referenced_object_id
JOIN sys.schemas AS sr ON sr.schema_id = tr.schema_id
JOIN sys.columns AS cr ON cr.object_id = fkc.referenced_object_id AND cr.column_id = fkc.referenced_column_id
ORDER BY fk.name, fkc.constraint_column_id;
"""
SQL_CHECK = """
SELECT s.name AS esquema, t.name AS tabla, cc.name AS restriccion, cc.definition
FROM sys.check_constraints AS cc
JOIN sys.tables AS t ON t.object_id = cc.parent_object_id
JOIN sys.schemas AS s ON s.schema_id = t.schema_id
ORDER BY s.name, t.name, cc.name;
"""
SQL_FILAS = """
SELECT s.name AS esquema, t.name AS tabla, SUM(p.rows) AS filas
FROM sys.tables AS t JOIN sys.schemas AS s ON s.schema_id = t.schema_id
JOIN sys.partitions AS p ON p.object_id = t.object_id AND p.index_id IN (0, 1)
GROUP BY s.name, t.name;
"""


def _consulta(cursor, sql):
    columnas, filas = ejecutar(cursor, sql)[0]
    return [dict(zip(columnas, f)) for f in filas]


def _tipo(c: dict) -> str:
    t = c["tipo"]
    if t in ("varchar", "char", "varbinary", "binary"):
        return f"{t}({'max' if c['max_length'] == -1 else c['max_length']})"
    if t in ("nvarchar", "nchar"):
        return f"{t}({'max' if c['max_length'] == -1 else c['max_length'] // 2})"
    if t in ("decimal", "numeric"):
        return f"{t}({c['precision']},{c['scale']})"
    if t in ("datetime2", "datetimeoffset", "time"):
        return f"{t}({c['scale']})"
    return t


def _limpiar(expr: str | None) -> str:
    if not expr:
        return ""
    return re.sub(r"\s+", " ", expr.replace("[", "").replace("]", "")).strip()


def leer_definiciones() -> dict:
    definiciones = {}
    with ARCHIVO_DEFINICIONES.open(encoding="utf-8") as f:
        for fila in csv.DictReader(f, delimiter=";"):
            definiciones[(fila["esquema"], fila["tabla"], fila["columna"])] = (fila["definicion"], fila["dominio"])
    return definiciones


def leer_catalogo(cfg: ConfigBD) -> dict:
    with conectar(cfg, autocommit=True) as conexion:
        cur = conexion.cursor()
        return {"columnas": _consulta(cur, SQL_COLUMNAS), "indices": _consulta(cur, SQL_INDICES), "fk": _consulta(cur, SQL_FK),
                "check": _consulta(cur, SQL_CHECK), "filas": _consulta(cur, SQL_FILAS)}


def construir_diccionario(catalogo: dict, definiciones: dict) -> tuple[list[dict], dict]:
    claves = defaultdict(list)          # (esquema, tabla, columna) -> etiquetas
    for i in catalogo["indices"]:
        clave = (i["esquema"], i["tabla"], i["columna"])
        if i["is_primary_key"]:
            claves[clave].append("PK")
        elif i["is_unique_constraint"]:
            claves[clave].append(f"UQ {i['indice']}")
        else:
            claves[clave].append(f"UX {i['indice']}" + (f" (filtrado: {_limpiar(i['filter_definition'])})" if i["filter_definition"] else ""))
    for fk in catalogo["fk"]:
        claves[(fk["esquema"], fk["tabla"], fk["columna"])].append(f"FK → {fk['esquema_ref']}.{fk['tabla_ref']}.{fk['columna_ref']}")
    checks_columna = defaultdict(list)
    checks_tabla = defaultdict(list)
    for ck in catalogo["check"]:
        checks_tabla[(ck["esquema"], ck["tabla"])].append((ck["restriccion"], _limpiar(ck["definition"])))
        for col in set(re.findall(r"\[(\w+)\]", ck["definition"])):
            checks_columna[(ck["esquema"], ck["tabla"], col)].append(ck["restriccion"])
    filas = {(f["esquema"], f["tabla"]): f["filas"] for f in catalogo["filas"]}

    registros = []
    for c in catalogo["columnas"]:
        if c["esquema"] in ESQUEMAS_EXCLUIDOS:
            continue
        clave = (c["esquema"], c["tabla"], c["columna"])
        definicion, dominio = definiciones.get(clave, ("", ""))
        notas = []
        if c["is_identity"]:
            notas.append("IDENTITY")
        if c["is_computed"]:
            notas.append(f"calculada{' persistida' if c['is_persisted'] else ''}: {_limpiar(c['calculo'])}")
        if c["masking_function"]:
            notas.append(f"enmascarada: {c['masking_function']}")
        registros.append({
            "esquema": c["esquema"], "tabla": c["tabla"], "orden": c["column_id"], "columna": c["columna"],
            "tipo": _tipo(c), "nulo": "Sí" if c["is_nullable"] else "No",
            "predeterminado": _limpiar(c["predeterminado"]), "claves": "; ".join(claves.get(clave, [])),
            "restricciones_check": ", ".join(sorted(checks_columna.get(clave, []))),
            "notas": "; ".join(notas), "definicion": definicion, "dominio": dominio,
            "filas_tabla": filas.get((c["esquema"], c["tabla"]), 0)})
    return registros, {"checks_tabla": checks_tabla, "filas": filas}


def _md(texto: str) -> str:
    return str(texto).replace("|", "\\|").replace("\n", " ")


def escribir_markdown(registros: list[dict], extra: dict, definiciones: dict) -> Path:
    por_tabla = defaultdict(list)
    for r in registros:
        por_tabla[(r["esquema"], r["tabla"])].append(r)
    total_columnas = len(registros)
    lineas = ["# Diccionario de datos", "",
              "Generado por `python -m aerolinea documentar` a partir del catálogo de SQL Server (tipos, nulabilidad, "
              "predeterminados, claves, restricciones y enmascaramiento) y de `docs/02_modelo/definiciones_atributos.csv` "
              "(definición y dominio de negocio). No editar a mano.", "",
              f"Tablas: {len(por_tabla)} · Columnas: {total_columnas} · Esquema `stg` (área de carga) excluido.", "",
              "Convenciones: **PK** clave primaria · **FK** clave foránea · **UQ** restricción única · **UX** índice único "
              "(filtrado cuando se indica) · CHECK = restricciones de comprobación que mencionan la columna.", ""]
    esquema_actual = None
    for (esquema, tabla), cols in sorted(por_tabla.items()):
        if esquema != esquema_actual:
            lineas += [f"## Esquema `{esquema}`", ""]
            esquema_actual = esquema
        descripcion = definiciones.get((esquema, tabla, ""), ("", ""))[0]
        lineas += [f"### `{esquema}.{tabla}`", "", f"{descripcion}  ", f"Filas al generar el documento: {extra['filas'].get((esquema, tabla), 0):,}".replace(",", "."), "",
                   "| Columna | Tipo | Nulo | Predeterminado | Claves | CHECK | Notas | Definición | Dominio |",
                   "|---|---|---|---|---|---|---|---|---|"]
        for r in sorted(cols, key=lambda x: x["orden"]):
            lineas.append("| " + " | ".join(_md(r[k]) for k in ("columna", "tipo", "nulo", "predeterminado", "claves",
                                                                   "restricciones_check", "notas", "definicion", "dominio")) + " |")
        checks = extra["checks_tabla"].get((esquema, tabla), [])
        if checks:
            lineas += ["", "Restricciones CHECK:", ""]
            lineas += [f"- `{nombre}`: `{_md(definicion)}`" for nombre, definicion in checks]
        lineas.append("")
    ruta = DIR_MODELO / "diccionario_datos.md"
    ruta.write_text("\n".join(lineas), encoding="utf-8")
    return ruta


def escribir_tabulares(registros: list[dict]) -> None:
    campos = ["esquema", "tabla", "orden", "columna", "tipo", "nulo", "predeterminado", "claves", "restricciones_check",
              "notas", "definicion", "dominio", "filas_tabla"]
    with (DIR_MODELO / "diccionario_datos.csv").open("w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.DictWriter(f, fieldnames=campos, delimiter=";")
        escritor.writeheader()
        escritor.writerows(registros)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Diccionario"
    hoja.append(campos)
    for r in registros:
        hoja.append([r[c] for c in campos])
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="256ABF")
    anchos = {"A": 8, "B": 24, "C": 6, "D": 26, "E": 18, "F": 6, "G": 22, "H": 40, "I": 30, "J": 30, "K": 50, "L": 40, "M": 10}
    for col, ancho in anchos.items():
        hoja.column_dimensions[col].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.freeze_panes = "E2"
    hoja.auto_filter.ref = hoja.dimensions
    libro.save(DIR_MODELO / "diccionario_datos.xlsx")


def _texto_sql(texto: str) -> str:
    return "N'" + texto.replace("'", "''") + "'"


def escribir_script_descripciones(registros: list[dict], definiciones: dict) -> Path:
    lineas = ["/* =============================================================================",
              "   13_descripciones.sql  (ARCHIVO GENERADO por python -m aerolinea documentar)",
              "   Propósito : registrar la definición de tablas y columnas como propiedades",
              "               extendidas MS_Description, visibles en SSMS y en Azure SQL.",
              "   Reinicio  : idempotente (actualiza si existe, agrega si no).",
              "   Depende de: 02 a 12 (las tablas deben existir).",
              "   ============================================================================= */",
              "SET NOCOUNT ON;",
              "IF DB_NAME() IN (N'master', N'model', N'msdb', N'tempdb')",
              "BEGIN",
              "    RAISERROR(N'Conéctese a la base del proyecto antes de ejecutar este script.', 16, 1);",
              "    SET NOEXEC ON;",
              "END",
              "GO"]
    tablas = sorted({(r["esquema"], r["tabla"]) for r in registros})
    objetivos = [(e, t, "", definiciones.get((e, t, ""), ("", ""))[0]) for e, t in tablas]
    objetivos += [(r["esquema"], r["tabla"], r["columna"], r["definicion"]) for r in registros if r["definicion"]]
    for esquema, tabla, columna, texto in objetivos:
        if not texto:
            continue
        nivel2 = f", @level2type = N'COLUMN', @level2name = N'{columna}'" if columna else ""
        filtro = f"N'{esquema}', N'TABLE', N'{tabla}', " + (f"N'COLUMN', N'{columna}'" if columna else "NULL, NULL")
        lineas.append(
            f"IF EXISTS (SELECT 1 FROM sys.fn_listextendedproperty(N'MS_Description', N'SCHEMA', {filtro})) "
            f"EXEC sys.sp_updateextendedproperty @name = N'MS_Description', @value = {_texto_sql(texto)}, @level0type = N'SCHEMA', "
            f"@level0name = N'{esquema}', @level1type = N'TABLE', @level1name = N'{tabla}'{nivel2}; "
            f"ELSE EXEC sys.sp_addextendedproperty @name = N'MS_Description', @value = {_texto_sql(texto)}, @level0type = N'SCHEMA', "
            f"@level0name = N'{esquema}', @level1type = N'TABLE', @level1name = N'{tabla}'{nivel2};")
    lineas += ["GO", "IF NOT EXISTS (SELECT 1 FROM adm.registro_despliegue WHERE script = '13_descripciones')",
               "    INSERT adm.registro_despliegue (script, descripcion) VALUES ('13_descripciones', N'Descripciones MS_Description');", "GO",
               "SET NOEXEC OFF;", "GO"]
    SCRIPT_DESCRIPCIONES.write_text("\n".join(lineas) + "\n", encoding="utf-8")
    return SCRIPT_DESCRIPCIONES


def _entidad(nombre: str) -> str:
    return nombre.replace(".", "_")


def escribir_diagramas_fisicos(registros: list[dict], catalogo: dict) -> list[Path]:
    por_tabla = defaultdict(list)
    for r in registros:
        por_tabla[f"{r['esquema']}.{r['tabla']}"].append(r)
    pk_tabla = defaultdict(set)
    for i in catalogo["indices"]:
        if i["is_primary_key"]:
            pk_tabla[f"{i['esquema']}.{i['tabla']}"].add(i["columna"])
    relaciones = defaultdict(lambda: {"cols": [], "nulable": False})
    for fk in catalogo["fk"]:
        clave = (fk["fk"], f"{fk['esquema']}.{fk['tabla']}", f"{fk['esquema_ref']}.{fk['tabla_ref']}")
        relaciones[clave]["cols"].append(fk["columna"])
        relaciones[clave]["nulable"] |= bool(fk["is_nullable"])
    rutas = []
    for archivo, (titulo, tablas) in AREAS.items():
        lineas = ["---", f"title: Modelo físico - {titulo}", "---", "erDiagram"]
        for tabla in tablas:
            lineas.append(f"    {_entidad(tabla)} {{")
            for r in sorted(por_tabla[tabla], key=lambda x: x["orden"]):
                marcas = []
                if "PK" in r["claves"]:
                    marcas.append("PK")
                if "FK" in r["claves"]:
                    marcas.append("FK")
                if ("UQ " in r["claves"] or "UX " in r["claves"]) and "PK" not in marcas:
                    marcas.append("UK")
                nota = "NULL" if r["nulo"] == "Sí" else "NOT NULL"
                if r["notas"].startswith("calculada"):
                    nota = "calculada"
                lineas.append(f"        {r['tipo'].replace(' ', '')} {r['columna']} {','.join(marcas)} \"{nota}\"".replace("  ", " "))
            lineas.append("    }")
        for (nombre, hija, padre), datos in sorted(relaciones.items()):
            if hija not in tablas or padre not in tablas:
                continue
            uno_a_uno = set(datos["cols"]) == pk_tabla[hija]
            lado_hijo = "o|" if uno_a_uno else "o{"
            lado_padre = "|o" if datos["nulable"] else "||"
            lineas.append(f"    {_entidad(padre)} {lado_padre}--{lado_hijo} {_entidad(hija)} : \"{nombre}\"")
        ruta = DIR_DIAGRAMAS / f"{archivo}.mmd"
        ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")
        rutas.append(ruta)
    return rutas


def renderizar(rutas: list[Path], eco=print) -> None:
    """Exporta PNG, SVG y PDF con Graphviz (diagramas compactos, legibles en papel).

    Si Graphviz no está instalado, recurre a mermaid-cli (variable MMDC o 'mmdc' en el PATH).
    """
    import os
    if shutil.which("dot"):
        from .diagramas import exportar
        for ruta in rutas:
            exportar(ruta, eco=eco)
        return
    mmdc = os.environ.get("MMDC") or shutil.which("mmdc")
    if not mmdc:
        eco("  mermaid-cli (mmdc) no está disponible: se conservan los .mmd (GitHub los muestra; draw.io los importa).")
        return
    extra = ["-p", os.environ["MMDC_PUPPETEER"]] if os.environ.get("MMDC_PUPPETEER") else []
    for ruta in rutas:
        for formato, opciones in (("svg", []), ("png", ["-s", "2"])):
            subprocess.run([mmdc, *extra, "-i", str(ruta), "-o", str(ruta.with_suffix(f".{formato}")), "-b", "white", *opciones],
                           check=True, capture_output=True)
        eco(f"  {ruta.name} -> .svg, .png")


def documentar(cfg: ConfigBD, eco=print) -> None:
    definiciones = leer_definiciones()
    catalogo = leer_catalogo(cfg)
    registros, extra = construir_diccionario(catalogo, definiciones)
    faltantes = [f"{r['esquema']}.{r['tabla']}.{r['columna']}" for r in registros if not r["definicion"]]
    eco(f"  Diccionario: {len(registros)} columnas; sin definición: {len(faltantes)} {faltantes[:10]}")
    escribir_markdown(registros, extra, definiciones)
    escribir_tabulares(registros)
    escribir_script_descripciones(registros, definiciones)
    rutas = escribir_diagramas_fisicos(registros, catalogo)
    renderizar(rutas + sorted(DIR_DIAGRAMAS.glob("conceptual*.mmd")) + sorted(DIR_DIAGRAMAS.glob("logico*.mmd")), eco)
